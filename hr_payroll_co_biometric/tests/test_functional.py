import base64
import io
from datetime import date, datetime, timedelta
from types import SimpleNamespace
from unittest.mock import MagicMock, patch

import openpyxl
import pytz

from odoo import fields
from odoo.exceptions import AccessError, UserError, ValidationError
from odoo.tests import TransactionCase, tagged, new_test_user
from odoo.addons.hr_payroll_co_biometric.hooks import post_init_hook


@tagged('post_install', '-at_install')
class TestFunctional(TransactionCase):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.env.company.partner_id.tz = 'America/Bogota'
        cls.partner = cls.env['res.partner'].create({'name': 'Functional Employee', 'vat': '991231234'})
        cls.employee = cls.env['hr.employee'].create({'partner_id': cls.partner.id, 'device_id_num': '99123'})
        cls.device = cls.env['biometric.device.details'].create({
            'name': 'Simulated ZK', 'device_ip': '192.0.2.1', 'device_timezone': 'America/Bogota'})
        cls.shift = cls.env['hr.shift'].create({
            'name': 'Functional Day', 'code': 'F19', 'hour_from': 8, 'hour_to': 16, 'rest_time': 0})
        cls.types = {}
        for code, rate in [('HED', 1.25), ('HEN', 1.75), ('RN', .35), ('HEDF', 2.0)]:
            cls.types[code] = cls.env['hr.overtime.type'].create({
                'name': code, 'code': code, 'category': 'earnings', 'rate': rate})
        post_init_hook(cls.env)
        cls.schedule = cls.env['hr.shift.schedule'].create({'date_start': '2026-09-21', 'date_end': '2026-09-30'})
        variable = cls.env['economic.variable'].search([('code', '=', 'MAX_HOURS_WORK')], limit=1)
        cls.env['economic.variable.line'].create({'variable_id': variable.id, 'year': '2026', 'month': '1', 'value': 240})
        cls.contract = cls.env['hr.contract'].create({
            'employee_id': cls.employee.id, 'state': 'open', 'date_start': '2026-01-01',
            'wage': 2400000, 'resource_calendar_id': cls.env.company.resource_calendar_id.id})

        cls.env['wage.update.history'].create({'contract_id': cls.contract.id, 'date': '2026-01-01', 'wage': 2400000})

    def setUp(self):
        super().setUp()
        # Device jobs explicitly commit; keep fixtures inside the test transaction.
        self.startPatcher(patch.object(self.env.cr, 'commit', return_value=None))

    def line(self, day='2026-09-24', shift=None):
        return self.env['hr.shift.schedule.line'].create({
            'schedule_id': self.schedule.id, 'employee_id': self.employee.id,
            'shift_id': (shift or self.shift).id, 'work_date': day})

    def punch(self, *times):
        result = self.env['biometric.raw.punch'].create([
            {'employee_id': self.employee.id, 'biometric_id': '99123',
             'device_id': self.device.id, 'punch_time': t} for t in times])
        result.action_process_imported_punches()
        self.env.flush_all()
        return result

    def record(self, start='2026-09-24 13:00:00', end='2026-09-24 21:00:00'):
        times = [t for t in (start, end) if t]
        if times:
            self.punch(*times)
        line = self.line()
        line.action_process_attendance()
        return line.attendance_record_id

    def wizard(self, **kwargs):
        values = {'date_start': '2026-09-24', 'date_end': '2026-09-24', 'schedule_id': self.schedule.id,
                  'employee_ids': [fields.Command.set(self.employee.ids)]}
        values.update(kwargs)
        return self.env['hr.process.attendance.wizard'].create(values)

    def workbook(self, rows):
        wb = openpyxl.Workbook()
        wb.active.append(['ID', 'Shift', 'Date'])
        for row in rows:
            wb.active.append(row)
        buf = io.BytesIO(); wb.save(buf)
        return base64.b64encode(buf.getvalue())

    def importer(self, rows, **values):
        return self.env['hr.import.shifts.excel.wizard'].create({
            'schedule_id': self.schedule.id, 'excel_file': self.workbook(rows), **values})

    def reporter(self, **values):
        return self.env['hr.biometric.overtime.report.wizard'].create({
            'date_start': '2026-09-21', 'date_end': '2026-09-30',
            'employee_ids': [fields.Command.set(self.employee.ids)], **values})

    def connection(self):
        conn = MagicMock()
        conn.get_users.return_value = [SimpleNamespace(uid=1, user_id='99123', name='Functional Employee')]
        conn.get_attendance.return_value = []
        zk = MagicMock(); zk.connect.return_value = conn
        return zk, conn

    def test_schedule_lifecycle(self):
        with self.assertRaises(UserError): self.schedule.action_confirm()
        self.line()
        self.schedule.action_confirm(); self.assertEqual(self.schedule.state, 'confirmed')
        self.assertEqual(self.schedule.action_process_attendance()['res_model'], 'hr.process.attendance.wizard')
        self.wizard().action_execute()
        self.assertEqual(self.schedule.line_ids.state, 'absent')
        self.assertEqual(self.schedule.attendance_record_count, 1)
        self.assertEqual(self.employee.attendance_record_count, 1)
        self.assertEqual(self.employee.action_view_attendance_records()['res_model'], 'hr.attendance.record')
        self.assertEqual(self.schedule.action_view_attendance_records()['res_model'], 'hr.attendance.record')
        self.schedule.action_done(); self.assertEqual(self.schedule.state, 'done')
        self.schedule.action_cancel(); self.assertFalse(self.schedule.line_ids.attendance_record_id)
        self.schedule.action_draft(); self.assertEqual(self.schedule.state, 'draft')

    def test_schedule_validation(self):
        with self.assertRaises(ValidationError), self.cr.savepoint():
            self.schedule.write({'date_end': '2026-09-01'})
        self.line()
        with self.assertRaises(ValidationError), self.cr.savepoint(): self.line()
        with self.assertRaises(ValidationError), self.cr.savepoint():
            self.wizard(date_end='2026-09-01')
        with self.assertRaises(UserError): self.schedule.action_process_attendance()

    def test_mass_generation_and_repeat(self):
        w = self.wizard(action='generate', shift_id=self.shift.id, date_end='2026-09-27')
        w.action_execute(); self.assertEqual(len(self.schedule.line_ids), 3)
        w.action_execute(); self.assertEqual(len(self.schedule.line_ids), 3)
        self.assertIn('Omitidas (ya existían): 3', w.result_message)

    def test_generation_by_employee(self):
        w = self.wizard(action='generate', shift_id=self.shift.id, shift_assignment_mode='by_employee')
        w._onchange_shift_assignment_mode()
        self.assertEqual(w.employee_shift_line_ids.employee_id, self.employee)
        w.action_execute(); self.assertEqual(self.schedule.line_ids.shift_id, self.shift)

    def test_weekday_schedule(self):
        self.shift.write({'assignment_mode': 'weekdays', 'weekday_line_ids': [fields.Command.create({
            'weekday': '3', 'hour_from': 9, 'hour_to': 15})]})
        self.assertTrue(self.shift.applies_on_date(date(2026, 9, 24)))
        self.assertFalse(self.shift.applies_on_date(date(2026, 9, 25)))
        self.assertEqual(self.line().planned_start, datetime(2026, 9, 24, 14))
        with self.assertRaises(ValidationError), self.cr.savepoint(): self.line('2026-09-25')

    def test_attendance_absence(self):
        rec = self.record(False, False)
        self.assertEqual(rec.attendance_result, 'total_absence')
        self.assertEqual(rec.novelty_ids.novelty_type, 'absence')

    def test_attendance_single_punch(self):
        self.assertEqual(self.record(end=False).attendance_result, 'biometric_issue')

    def test_attendance_minor_delay(self):
        rec = self.record(start='2026-09-24 13:15:00')
        self.assertEqual(rec.attendance_result, 'late_minor')
        self.assertEqual(rec.novelty_ids.minutes, 15)

    def test_attendance_major_delay(self):
        self.assertEqual(self.record(start='2026-09-24 13:45:00').attendance_result, 'late_major')

    def test_attendance_early_departure(self):
        self.assertEqual(self.record(end='2026-09-24 19:00:00').attendance_result, 'partial_absence')

    def test_overtime_payroll_and_immutability(self):
        rec = self.record(end='2026-09-24 23:00:00')
        self.assertEqual(rec.attendance_result, 'extra_hours')
        self.assertAlmostEqual(rec.extra_day_hours, 2)
        novelty = rec.novelty_ids.filtered(lambda n: n.novelty_type == 'extra_day_hours')
        self.assertEqual(len(novelty), 1)
        novelty.action_to_review(); self.assertEqual(novelty.state, 'review')
        novelty.action_approve(); novelty.action_apply()
        self.assertEqual(novelty.overtime_id.contract_id, self.contract)
        self.assertEqual(novelty.overtime_id.qty, 2)
        self.assertEqual(novelty.overtime_id.overtime_type_id, self.types['HED'])
        with self.assertRaises(UserError): novelty.write({'extra_hours': 3})
        with self.assertRaises(UserError): novelty.unlink()
        with self.assertRaises(UserError): self.schedule.action_cancel()
        before = novelty.overtime_id
        rec.action_compute_from_biometric()
        self.assertEqual(novelty.overtime_id, before)
        self.assertEqual(len(rec.novelty_ids.filtered(lambda n: n.novelty_type == 'extra_day_hours')), 1)

    def test_overtime_day_night_split(self):
        rec = self.record(end='2026-09-25 01:00:00')
        self.assertAlmostEqual(rec.extra_day_hours, 3)
        self.assertAlmostEqual(rec.extra_night_hours, 1)
        self.assertEqual(set(rec.novelty_ids.mapped('overtime_type_id.code')), {'HED', 'HEN'})

    def test_tolerance_setting(self):
        settings = self.env['res.config.settings'].create({'overtime_tolerance_minutes': 30})
        self.assertEqual(self.env.company.overtime_tolerance_minutes, 30)
        rec = self.record(end='2026-09-24 21:20:00')
        self.assertEqual(rec.extra_hours, 0)
        self.assertEqual(rec.attendance_result, 'ok')

    def test_overtime_limits(self):
        rec = self.record(end='2026-09-25 01:00:00')
        self.assertTrue(rec.overtime_limit_alert)
        w = self.wizard(action='check_overtime'); w.action_execute()
        self.assertTrue(w.has_overtime_alert)
        self.assertIn('Límite diario', w.overtime_alert_message)

    def test_correction_workflow(self):
        rec = self.record()
        rec.action_request_correction(); self.assertEqual(rec.correction_state, 'requested')
        rec.action_reject_correction(); self.assertEqual(rec.correction_state, 'rejected')
        rec.action_request_correction(); rec.action_approve_correction()
        self.assertEqual(rec.correction_state, 'approved')
        self.assertEqual(rec.correction_approved_by, self.env.user)
        self.assertEqual(rec.action_view_novelties()['view_mode'], 'list,form')

    def test_novelty_rejection_and_mass_apply(self):
        rec = self.record(False, False); nov = rec.novelty_ids
        with self.assertRaises(UserError): nov.action_apply()
        nov.action_to_review(); nov.action_reject(); self.assertEqual(nov.state, 'rejected')
        other = nov.copy({'state': 'draft'})
        other.action_approve(); other.action_apply_mass(); self.assertEqual(other.state, 'applied')
        with self.assertRaises(UserError): (nov | other).action_apply_mass()

    def test_rules_hook_and_invalid_ranges(self):
        rules = self.env['hr.overtime.rule'].search([('company_id', '=', self.env.company.id)])
        post_init_hook(self.env)
        self.assertEqual(self.env['hr.overtime.rule'].search_count([('company_id', '=', self.env.company.id)]), len(rules))
        with self.assertRaises(ValidationError), self.cr.savepoint(): rules[0].write({'hour_from': 25})
        with self.assertRaises(ValidationError), self.cr.savepoint(): rules[0].write({'hour_from': 8, 'hour_to': 8})

    def test_excel_import_valid_and_invalid_rows(self):
        w = self.importer([['991231234', 'F19', '2026-09-24'], ['missing', 'F19', '2026-09-25'],
                           ['991231234', 'bad', '2026-09-25'], ['991231234', 'F19', 'bad']])
        w.action_import(); self.assertEqual(len(self.schedule.line_ids), 1)
        self.assertIn('Errores (3)', w.result_html)
        self.assertEqual(w.action_open_schedule()['res_id'], self.schedule.id)
        w.action_import(); self.assertEqual(len(self.schedule.line_ids), 1)

    def test_excel_duplicate_rows_in_same_file(self):
        row = ['991231234', 'F19', '2026-09-24']
        w = self.importer([row, row]); w.action_import()
        self.assertEqual(len(self.schedule.line_ids), 1)
        self.assertIn('Duplicados omitidos:</strong> 1', w.result_html)

    def test_excel_template_and_new_schedule(self):
        w = self.importer([['991231234', 'F19', '24/09/2026']], create_new_schedule=True,
                          schedule_name='New XLSX', date_start='2026-09-24', date_end='2026-09-24')
        action = w.action_download_template()
        attachment = self.env['ir.attachment'].browse(int(action['url'].split('/')[3].split('?')[0]))
        wb = openpyxl.load_workbook(io.BytesIO(base64.b64decode(attachment.datas)))
        self.assertEqual(wb.sheetnames, ['Plantilla Turnos', 'Turnos Disponibles'])
        w.action_import()
        self.assertNotEqual(w.action_open_schedule()['res_id'], self.schedule.id)

    def test_excel_invalid_file_and_confirmed_schedule(self):
        w = self.importer([]); w.excel_file = base64.b64encode(b'not xlsx')
        with self.assertRaises(UserError): w.action_import()
        self.line(); self.schedule.action_confirm()
        w.excel_file = self.workbook([['991231234', 'F19', '2026-09-25']])
        with self.assertRaises(UserError): w.action_import()

    def test_report_excel_and_qweb(self):
        self.record(); w = self.reporter()
        self.assertEqual(w.get_report_data()['total_day'], 8)
        action = w.action_export_excel()
        attachment = self.env['ir.attachment'].browse(int(action['url'].split('/')[3].split('?')[0]))
        wb = openpyxl.load_workbook(io.BytesIO(base64.b64decode(attachment.datas)), data_only=True)
        self.assertTrue(any('Functional Employee' in str(c.value) for row in wb.active for c in row))
        self.assertEqual(w.with_context(discard_logo_check=True).action_print_pdf()['report_type'], 'qweb-pdf')
        report = self.env.ref('hr_payroll_co_biometric.action_report_biometric_overtime')
        html, _ = report._render_qweb_html(report.report_name, w.ids, data={'wizard_id': w.id})
        self.assertIn(b'Functional Employee', html)

    def test_report_empty_filter_does_not_return_other_employees(self):
        self.record()
        department = self.env['hr.department'].create({'name': 'Empty department'})
        w = self.reporter(department_id=department.id)
        self.assertFalse(w._get_records())
        with self.assertRaises(UserError): w.action_export_excel()
        with self.assertRaises(UserError): w.action_print_pdf()

    def test_raw_punch_delete_preserves_shared_log(self):
        punches = self.punch('2026-09-24 13:00:00')
        duplicate = punches.copy(); duplicate.action_process_imported_punches()
        log = punches.zk_attendance_id
        punches.unlink(); self.assertTrue(log.exists())
        duplicate.unlink(); self.assertFalse(log.exists())

    def test_raw_punch_unknown_employee(self):
        raw = self.env['biometric.raw.punch'].create({'biometric_id': 'unknown', 'device_id': self.device.id,
                                                   'punch_time': '2026-09-24 13:00:00'})
        self.assertEqual(raw.action_process_imported_punches()['params']['type'], 'warning')
        self.assertFalse(raw.zk_attendance_id)

    def test_daily_analysis_and_native_attendance(self):
        self.punch('2026-09-24 13:00:00', '2026-09-24 17:00:00', '2026-09-24 21:00:00')
        daily = self.env['daily.attendance'].search([('employee_id', '=', self.employee.id)])
        self.assertEqual(daily.punch_count, 3); self.assertEqual(daily.ignored_count, 1)
        self.assertEqual(daily.ignored_punches, '12:00:00')
        with self.assertRaises(UserError): daily.unlink()
        with patch('odoo.fields.Date.today', return_value=date(2026, 9, 25)):
            self.device._sync_to_hr_attendance(); self.device._sync_to_hr_attendance()
        native = self.env['hr.attendance'].search([('employee_id', '=', self.employee.id)])
        self.assertEqual(len(native), 1)
        self.assertEqual(native.check_in, datetime(2026, 9, 24, 13))
        self.assertEqual(native.check_out, datetime(2026, 9, 24, 21))

    def test_device_connection_and_disconnect(self):
        zk, conn = self.connection()
        with patch.object(type(self.device), '_get_zk_connection', return_value=zk):
            self.assertEqual(self.device.action_test_connection()['params']['type'], 'success')
        conn.disconnect.assert_called_once()
        zk.connect.side_effect = OSError('offline')
        with patch.object(type(self.device), '_get_zk_connection', return_value=zk):
            with self.assertRaises(ValidationError): self.device.action_test_connection()
            self.assertFalse(self.device._test_connection_logic()[0])
        self.assertEqual(self.device.status, 'error')

    def test_device_time_restart_clear_and_crons(self):
        zk, conn = self.connection()
        with patch.object(type(self.device), '_get_zk_connection', return_value=zk):
            self.device.action_set_timezone(); self.device.cron_sync_time()
            self.assertEqual(conn.set_time.call_count, 2)
            self.device.action_restart_device(); conn.restart.assert_called_once()
            self.device.action_clear_attendance_device(); conn.clear_attendance.assert_called_once()
            self.device.cron_connection_test(); self.assertEqual(self.device.status, 'ok')

    def test_device_download_timezone_deduplication(self):
        zk, conn = self.connection()
        conn.get_attendance.return_value = [SimpleNamespace(user_id='099123', timestamp=datetime(2026, 9, 24, h),
                                                          punch=0, status=1) for h in (8, 16)]
        with patch.object(type(self.device), '_get_zk_connection', return_value=zk):
            self.device.cron_download(); self.device.cron_download()
        raw = self.env['biometric.raw.punch'].search([('employee_id', '=', self.employee.id)])
        self.assertEqual(len(raw), 2)
        self.assertEqual(min(raw.mapped('punch_time')), datetime(2026, 9, 24, 13))
        self.assertEqual(self.env['zk.machine.attendance'].search_count([('employee_id', '=', self.employee.id)]), 2)
        self.assertEqual(self.device.status, 'ok')

    def test_device_download_failure_reenables_device(self):
        zk, conn = self.connection(); conn.get_attendance.side_effect = OSError('read failed')
        with patch.object(type(self.device), '_get_zk_connection', return_value=zk): self.device.action_download_attendance()
        self.assertEqual(self.device.status, 'error')
        conn.enable_device.assert_called_once(); conn.disconnect.assert_called_once()

    def test_device_whitelist(self):
        other = self.device.copy({'name': 'Allowed only'})
        self.employee.allowed_device_ids = other
        zk, conn = self.connection()
        conn.get_attendance.return_value = [SimpleNamespace(user_id='99123', timestamp=datetime(2026, 9, 24, 8), punch=0, status=1)]
        with patch.object(type(self.device), '_get_zk_connection', return_value=zk): self.device.action_download_attendance()
        self.assertFalse(self.env['zk.machine.attendance'].search([('employee_id', '=', self.employee.id)]))
        self.assertEqual(self.env['biometric.raw.punch'].search_count([('employee_id', '=', self.employee.id)]), 1)

    def test_device_user_mapping_upload_delete(self):
        zk, conn = self.connection()
        with patch.object(type(self.device), '_get_zk_connection', return_value=zk):
            action = self.device.action_open_user_mapping_wizard()
            self.assertEqual(action['res_model'], 'biometric.user.mapping')
            mapping = self.env['biometric.user.mapping'].search([('device_id', '=', self.device.id)])
            self.assertEqual(mapping.employee_id, self.employee)
            mapping.write({'employee_id': self.employee.id})
            w = self.env['biometric.odoo.to.biometric.wizard'].with_context(active_id=self.device.id).create({
                'employee_ids': [fields.Command.set(self.employee.ids)]})
            w.action_upload_users(); self.assertEqual(conn.set_user.call_args.kwargs['user_id'], '99123')
            conn.get_users.side_effect = [[SimpleNamespace(uid=1, user_id='99123', name='Test')], []]
            mapping.action_delete_user_from_device(); conn.delete_user.assert_called_once_with(uid=1)

    def test_access_read_only_and_import_roles(self):
        user = new_test_user(self.env, context={'no_reset_password': True}, login='biometric_reader19', groups='base.group_user,hr_payroll_co_biometric.group_biometric_user')
        self.assertTrue(self.shift.with_user(user).read(['name']))
        with self.assertRaises(AccessError): self.shift.with_user(user).write({'name': 'Forbidden'})
        raw = self.punch('2026-09-24 13:00:00')
        with self.assertRaises(AccessError): raw.with_user(user).action_process_imported_punches()
        importer = new_test_user(self.env, context={'no_reset_password': True}, login='biometric_importer19', groups='base.group_user,hr_zk_attendance.group_biometric_raw_punch_importer')
        self.assertEqual(raw.with_user(importer).action_process_imported_punches()['params']['type'], 'success')
        with self.assertRaises(AccessError): raw.with_user(importer).write({'biometric_id': 'other'})

    def test_shift_display_name_contains_code(self):
        self.assertIn('F19', self.shift.display_name)

    def test_sunday_and_holiday_overtime(self):
        self.punch('2026-09-27 13:00:00', '2026-09-27 21:00:00')
        line = self.line('2026-09-27'); self.assertTrue(line.is_sunday)
        line.action_process_attendance(); record = line.attendance_record_id
        self.assertEqual(record.holiday_work_hours, 8)
        novelty = record.novelty_ids.filtered(lambda n: n.novelty_type == 'compensatory')
        self.assertEqual(novelty.extra_hours, 8)
        novelty.action_approve(); novelty.action_apply()
        self.assertEqual(novelty.overtime_id.overtime_type_id, self.types['HEDF'])
        self.assertEqual(novelty.overtime_id.amount, 160000)

    def test_holiday_excluded_from_generation(self):
        year = self.env['hr.days.off.year'].create({'year': '2026'})
        self.env['hr.day.off'].create({'name': 'Test holiday', 'date': '2026-09-24', 'year_id': year.id})
        w = self.wizard(action='generate', shift_id=self.shift.id, date_end='2026-09-25')
        w.action_execute()
        self.assertEqual(self.schedule.line_ids.work_date, date(2026, 9, 25))
        self.assertTrue(self.line().is_holiday)

    def test_night_surcharge_with_rest(self):
        self.shift.write({'hour_from': 22, 'hour_to': 6, 'rest_time': .5})
        self.punch('2026-09-25 03:00:00', '2026-09-25 11:00:00')
        line = self.line(); line.action_process_attendance()
        self.assertEqual(self.shift.effective_hours, 7.5)
        record = line.attendance_record_id
        # Source behavior: attendance is the gross interval; rest only changes effective_hours.
        self.assertEqual(record.night_hours, 8)
        novelty = record.novelty_ids.filtered(lambda n: n.novelty_type == 'night_surcharge')
        novelty.action_approve(); novelty.action_apply()
        self.assertEqual(novelty.overtime_id.overtime_type_id, self.types['RN'])

    def test_zero_tolerance_and_disabled_shift_tolerance(self):
        self.shift.write({'tolerance_in': 0, 'tolerance_out': 0})
        rec = self.record(start='2026-09-24 13:01:00')
        self.assertEqual(rec.attendance_result, 'late_minor')
        self.shift.write({'tolerance_out': 30, 'apply_overtime_tolerance': False})
        rec.actual_check_out = datetime(2026, 9, 24, 21, 11)
        self.assertAlmostEqual(rec.extra_hours, 10 / 60)

    def test_settings_recompute_and_threshold(self):
        rec = self.record(end='2026-09-24 21:30:00')
        self.env.company.overtime_tolerance_minutes = 31
        self.assertEqual(rec.extra_hours, 0)
        self.env.company.overtime_tolerance_minutes = 30
        self.assertEqual(rec.extra_hours, .5)

    def test_weekly_overtime_limit(self):
        records = self.env['hr.attendance.record']
        for day in range(21, 26):
            self.punch('2026-09-%s 13:00:00' % day, '2026-09-%s 00:00:00' % (day + 1))
            line = self.line('2026-09-%s' % day); line.action_process_attendance()
            records |= line.attendance_record_id
        records.invalidate_recordset(['weekly_extra_hours', 'overtime_limit_alert'])
        self.assertTrue(all(r.weekly_extra_hours == 15 for r in records))
        w = self.wizard(action='check_overtime', date_start='2026-09-21', date_end='2026-09-25')
        w.action_execute(); self.assertIn('Límite semanal', w.overtime_alert_message)

    def test_daily_processing_cron(self):
        line = self.line(); self.schedule.action_confirm()
        with patch('odoo.fields.Date.today', return_value=date(2026, 9, 25)):
            self.env['hr.shift.schedule.line']._cron_process_daily_attendance()
        self.assertEqual(line.state, 'absent')

    def test_payroll_missing_contract_and_zero_rate(self):
        rec = self.record(end='2026-09-24 23:00:00'); nov = rec.novelty_ids
        nov.action_approve(); self.contract.state = 'close'
        with self.assertRaises(UserError): nov.action_apply()
        self.contract.state = 'open'; self.types['HED'].rate = 0
        with self.assertRaises(UserError): nov.action_apply()

    def test_economic_value_and_payroll_amount(self):
        rec = self.record(end='2026-09-24 23:00:00')
        self.assertEqual(rec._get_hour_base_value(), 10000)
        self.assertEqual(rec.total_overtime_value, 25000)
        rec.novelty_ids.action_approve(); rec.novelty_ids.action_apply()
        self.assertEqual(rec.novelty_ids.overtime_id.amount, 25000)

    def test_approved_novelty_is_not_duplicated_by_reprocessing(self):
        rec = self.record(end='2026-09-24 23:00:00')
        nov = rec.novelty_ids; nov.action_approve()
        rec.action_compute_from_biometric()
        self.assertEqual(rec.novelty_ids, nov)

    def test_report_pdf_binary(self):
        self.record(); w = self.reporter()
        report = self.env.ref('hr_payroll_co_biometric.action_report_biometric_overtime')
        content, fmt = report.with_context(force_report_rendering=True)._render_qweb_pdf(
            report.report_name, w.ids, data={'wizard_id': w.id})
        self.assertEqual(fmt, 'pdf'); self.assertTrue(content.startswith(b'%PDF'))

    def test_user_mapping_read_error_restores_device(self):
        zk, conn = self.connection(); conn.get_users.side_effect = OSError('failed')
        with patch.object(type(self.device), '_get_zk_connection', return_value=zk):
            with self.assertRaises(UserError): self.device.action_open_user_mapping_wizard()
        conn.enable_device.assert_called_once(); conn.disconnect.assert_called_once()

    def test_upload_missing_id_and_offline(self):
        w = self.env['biometric.odoo.to.biometric.wizard'].create({'device_id': self.device.id,
            'employee_ids_without_id': [fields.Command.set(self.employee.ids)]})
        self.employee.device_id_num = False
        with self.assertRaises(UserError): w.action_save_biometric_id()
        zk, conn = self.connection(); zk.connect.side_effect = OSError('offline')
        with patch.object(type(self.device), '_get_zk_connection', return_value=zk):
            with self.assertRaises(UserError): w.action_upload_users()
            with self.assertRaises(UserError): self.device.action_restart_device()
            with self.assertRaises(UserError): self.device.action_clear_attendance_device()
            with self.assertRaises(UserError): self.device.action_set_timezone()

    def test_clear_odoo_logs_on_isolated_transaction(self):
        self.punch('2026-09-24 13:00:00', '2026-09-24 21:00:00')
        self.device._sync_to_hr_attendance()
        self.device.action_clear_attendance_odoo()
        self.env.invalidate_all()
        self.assertFalse(self.env['zk.machine.attendance'].search([]))
        self.assertFalse(self.env['biometric.raw.punch'].search([]))
        self.assertFalse(self.env['hr.attendance'].search([]))

    def test_shift_constraints_and_rotating_extension_configuration(self):
        with self.assertRaises(ValidationError), self.cr.savepoint():
            self.shift.copy({'code': 'F19'})
        self.shift.write({'is_extension': True, 'extension_hours': 2,
                          'is_rotating_cycle': True, 'cycle_work_days': 2, 'cycle_rest_days': 2})
        self.assertEqual(self.shift.extension_hours, 2)
        self.assertTrue(self.shift.is_rotating_cycle)

    def test_documentation_rendering(self):
        from odoo.addons.hr_payroll_co_biometric.controllers.main import _md_to_html
        html = _md_to_html('# Title\n\n**Bold**\n\n| A | B |\n|---|---|\n| 1 | 2 |\n')
        self.assertIn('<h1>', html); self.assertIn('<strong>Bold</strong>', html)
        self.assertIn('<table', html)

    def test_all_backend_views_compile(self):
        views = self.env['ir.model.data'].search([
            ('module', 'in', ['hr_zk_attendance', 'hr_payroll_co_biometric']), ('model', '=', 'ir.ui.view')])
        checked = 0
        for xml in views:
            view = self.env['ir.ui.view'].browse(xml.res_id)
            if view.type == 'qweb': continue
            with self.subTest(view=xml.complete_name):
                result = self.env[view.model].get_view(view.id, view.type)
                self.assertTrue(result['arch']); checked += 1
        self.assertGreater(checked, 25)

    def test_role_acl_matrix(self):
        permissions = {
            'user': (False, False, False), 'auditor': (False, False, False),
            'supervisor': (False, True, False), 'hr': (False, True, True), 'manager': (True, True, True),
        }
        for role, expected in permissions.items():
            user = new_test_user(self.env, context={'no_reset_password': True}, login='acl_' + role,
                groups='base.group_user,hr_payroll_co_biometric.group_biometric_' + role)
            actual = tuple(self.env[model].with_user(user).has_access('create')
                           for model in ['hr.shift', 'hr.shift.schedule', 'hr.attendance.record'])
            self.assertEqual(actual, expected, role)

    def test_edit_users_menu_permissions_and_actions(self):
        user = new_test_user(self.env, context={'no_reset_password': True}, login='device_menu', groups='base.group_user,hr.group_hr_user')
        wizard = self.env['biometric.edit.users.menu.wizard'].with_user(user).with_context(active_id=self.device.id).create({})
        self.assertEqual(wizard.action_map_odoo_to_biometric()['res_model'], 'biometric.odoo.to.biometric.wizard')
        zk, conn = self.connection()
        with patch.object(type(self.device), '_get_zk_connection', return_value=zk):
            self.assertEqual(wizard.action_map_biometric_to_odoo()['res_model'], 'biometric.user.mapping')

    def test_known_issue_employee_role_reads_other_employee_attendance(self):
        """Known source issue: missing record rules; role is documented as own records only."""
        record = self.record()
        reader = new_test_user(self.env, context={'no_reset_password': True}, login='other_employee_reader',
            groups='base.group_user,hr_payroll_co_biometric.group_biometric_user')
        self.assertFalse(self.env['hr.attendance.record'].with_user(reader).search([('id', '=', record.id)]))

    def test_known_issue_supervisor_can_approve_hr_novelties(self):
        """Known source issue: write ACL also allows approval without HR role."""
        novelty = self.record(False, False).novelty_ids
        supervisor = new_test_user(self.env, context={'no_reset_password': True}, login='non_hr_supervisor',
            groups='base.group_user,hr_payroll_co_biometric.group_biometric_supervisor')
        with self.assertRaises(AccessError): novelty.with_user(supervisor).action_approve()

    def test_known_issue_internal_user_can_clear_all_attendance(self):
        """Known source issue: global SQL deletion has no manager check or company filter."""
        self.punch('2026-09-24 13:00:00')
        user = new_test_user(self.env, context={'no_reset_password': True}, login='non_manager_device_user', groups='base.group_user')
        with self.assertRaises(AccessError): self.device.with_user(user).action_clear_attendance_odoo()
