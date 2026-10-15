# -*- coding: utf-8 -*-
import base64
import io
from datetime import datetime, date

from odoo import models, fields, api
from odoo.exceptions import UserError, ValidationError


class HrImportShiftsExcelWizard(models.TransientModel):
    """Wizard para importar programación de turnos desde un archivo Excel.

    Formato de la plantilla (hoja 1):
    Columna A: Cédula o ID del empleado (obligatorio)
    Columna B: Código del turno (ej. TA, TB, ADM) (obligatorio)
    Columna C: Fecha (YYYY-MM-DD o DD/MM/YYYY) (obligatorio)
    """
    _name = 'hr.import.shifts.excel.wizard'
    _description = 'Importar Turnos desde Excel'

    # ─── Opción 1: Agregar a programación existente ────────────────────────
    schedule_id = fields.Many2one(
        'hr.shift.schedule',
        string='Programación Existente',
        domain="[('state', '=', 'draft')]",
        help='Si se selecciona, las líneas se agregan a esta programación (debe estar en borrador).',
    )

    # ─── Opción 2: Crear nueva programación ───────────────────────────────
    create_new_schedule = fields.Boolean(
        string='Crear Nueva Programación',
        default=False,
    )
    schedule_name = fields.Char(
        string='Nombre de la Programación',
        default='Importación desde Excel',
    )
    date_start = fields.Date(string='Fecha Inicio')
    date_end = fields.Date(string='Fecha Fin')

    # ─── Archivo ──────────────────────────────────────────────────────────
    excel_file = fields.Binary(
        string='Archivo Excel',
        required=True,
        help='Cargue el archivo .xlsx con el formato de la plantilla.',
    )
    excel_filename = fields.Char(string='Nombre del archivo')

    # ─── Resultado ────────────────────────────────────────────────────────
    result_html = fields.Html(string='Resultado', readonly=True)
    state = fields.Selection([
        ('upload', 'Cargar Archivo'),
        ('result', 'Resultado'),
    ], default='upload')

    @api.onchange('create_new_schedule')
    def _onchange_create_new_schedule(self):
        if self.create_new_schedule:
            self.schedule_id = False

    def action_download_template(self):
        """Genera y devuelve una plantilla Excel descargable."""
        try:
            import openpyxl
        except ImportError:
            raise UserError(
                'El módulo openpyxl no está instalado. '
                'Instálelo con: pip install openpyxl'
            )

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Plantilla Turnos'

        # Encabezados
        headers = ['Cédula/ID Empleado', 'Código Turno', 'Fecha (YYYY-MM-DD)']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = openpyxl.styles.Font(bold=True)
            cell.fill = openpyxl.styles.PatternFill(
                start_color='1F4E79', end_color='1F4E79', fill_type='solid'
            )
            cell.font = openpyxl.styles.Font(bold=True, color='FFFFFF')

        # Filas de ejemplo
        examples = [
            ['123456789', 'TA', '2026-06-01'],
            ['987654321', 'TB', '2026-06-01'],
            ['123456789', 'TC', '2026-06-02'],
        ]
        for row_idx, row_data in enumerate(examples, 2):
            for col_idx, value in enumerate(row_data, 1):
                ws.cell(row=row_idx, column=col_idx, value=value)

        # Ajustar anchos de columna
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            ws.column_dimensions[col[0].column_letter].width = max(max_len + 4, 20)

        # Segunda hoja: turnos disponibles
        ws2 = wb.create_sheet(title='Turnos Disponibles')
        ws2.cell(row=1, column=1, value='Código').font = openpyxl.styles.Font(bold=True)
        ws2.cell(row=1, column=2, value='Nombre').font = openpyxl.styles.Font(bold=True)
        shifts = self.env['hr.shift'].search([], order='code')
        for row_idx, shift in enumerate(shifts, 2):
            ws2.cell(row=row_idx, column=1, value=shift.code)
            ws2.cell(row=row_idx, column=2, value=shift.name)

        # Guardar en buffer
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        file_data = base64.b64encode(buffer.read()).decode()

        # Retornar como adjunto descargable
        attachment = self.env['ir.attachment'].create({
            'name': 'plantilla_turnos.xlsx',
            'type': 'binary',
            'datas': file_data,
            'mimetype': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        })
        return {
            'type': 'ir.actions.act_url',
            'url': '/web/content/%d?download=true' % attachment.id,
            'target': 'new',
        }

    def action_import(self):
        """Procesa el archivo Excel y crea las líneas de programación."""
        self.ensure_one()
        if not self.excel_file:
            raise UserError('Por favor cargue un archivo Excel.')

        try:
            import openpyxl
        except ImportError:
            raise UserError('El módulo openpyxl no está instalado. Instálelo con: pip install openpyxl')

        # Leer el archivo
        try:
            raw = base64.b64decode(self.excel_file)
            wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            ws = wb.active
        except Exception as e:
            raise UserError('Error al leer el archivo Excel: %s' % str(e))

        # Obtener o crear la programación destino
        if self.create_new_schedule:
            if not self.schedule_name:
                raise UserError('Ingrese un nombre para la nueva programación.')
            schedule = self.env['hr.shift.schedule'].create({
                'name': self.schedule_name,
                'date_start': self.date_start or fields.Date.today(),
                'date_end': self.date_end or fields.Date.today(),
            })
        elif self.schedule_id:
            schedule = self.schedule_id
            if schedule.state != 'draft':
                raise UserError('La programación seleccionada debe estar en estado Borrador.')
        else:
            raise UserError('Seleccione una programación existente o active la opción "Crear Nueva Programación".')

        # Cargar empleados y turnos en memoria para validación eficiente
        employees_by_id = {
            str(e.identification_id): e
            for e in self.env['hr.employee'].search([('active', '=', True)])
            if e.identification_id
        }
        shifts_by_code = {
            s.code.upper(): s
            for s in self.env['hr.shift'].search([])
            if s.code
        }

        errors = []
        created = 0
        duplicates = 0
        line_vals_list = []
        pending_assignments = set()

        for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if not row or not any(row[:3]):
                continue

            raw_id = str(row[0]).strip() if row[0] is not None else ''
            raw_code = str(row[1]).strip().upper() if row[1] is not None else ''
            raw_date = row[2]

            # Validar empleado
            if not raw_id:
                errors.append('Fila %d: Cédula vacía.' % row_num)
                continue
            employee = employees_by_id.get(raw_id)
            if not employee:
                errors.append('Fila %d: Empleado con cédula %s no encontrado.' % (row_num, raw_id))
                continue

            # Validar turno
            if not raw_code:
                errors.append('Fila %d: Código de turno vacío.' % row_num)
                continue
            shift = shifts_by_code.get(raw_code)
            if not shift:
                errors.append('Fila %d: Turno con código %s no encontrado.' % (row_num, raw_code))
                continue

            # Validar fecha
            work_date = None
            if isinstance(raw_date, (datetime,)):
                work_date = raw_date.date()
            elif isinstance(raw_date, date):
                work_date = raw_date
            elif isinstance(raw_date, str):
                for fmt in ('%Y-%m-%d', '%d/%m/%Y', '%d-%m-%Y'):
                    try:
                        work_date = datetime.strptime(raw_date.strip(), fmt).date()
                        break
                    except ValueError:
                        continue
            if not work_date:
                errors.append('Fila %d: Fecha inválida "%s". Use el formato YYYY-MM-DD.' % (row_num, raw_date))
                continue

            # Verificar duplicado en la misma programación
            existing = self.env['hr.shift.schedule.line'].search([
                ('schedule_id', '=', schedule.id),
                ('employee_id', '=', employee.id),
                ('work_date', '=', work_date),
            ], limit=1)
            assignment_key = (employee.id, work_date)
            if existing or assignment_key in pending_assignments:
                duplicates += 1
                continue

            pending_assignments.add(assignment_key)
            line_vals_list.append({
                'schedule_id': schedule.id,
                'employee_id': employee.id,
                'shift_id': shift.id,
                'work_date': work_date,
            })

        # Crear todas las líneas válidas
        if line_vals_list:
            self.env['hr.shift.schedule.line'].create(line_vals_list)
            created = len(line_vals_list)

        # Construir HTML de resultados
        summary_lines = [
            '<p><strong>Programación:</strong> %s</p>' % schedule.name,
            '<p><strong>Líneas creadas:</strong> %d</p>' % created,
            '<p><strong>Duplicados omitidos:</strong> %d</p>' % duplicates,
        ]
        if errors:
            error_items = ''.join('<li>%s</li>' % e for e in errors[:50])
            if len(errors) > 50:
                error_items += '<li>... y %d errores más.</li>' % (len(errors) - 50)
            summary_lines.append(
                '<div class="alert alert-warning mt-2">'
                '<strong>Errores (%d):</strong><ul>%s</ul></div>' % (len(errors), error_items)
            )
        else:
            summary_lines.append('<div class="alert alert-success mt-2">Importación completada sin errores.</div>')

        self.result_html = '\n'.join(summary_lines)
        self.schedule_id = schedule
        self.state = 'result'

        return {
            'type': 'ir.actions.act_window',
            'res_model': self._name,
            'res_id': self.id,
            'view_mode': 'form',
            'target': 'new',
            'name': 'Resultado de Importación',
        }

    def action_open_schedule(self):
        self.ensure_one()
        schedule = self.schedule_id or self.env['hr.shift.schedule'].search(
            [('name', '=', self.schedule_name)], order='id desc', limit=1
        )
        if not schedule:
            return {'type': 'ir.actions.act_window_close'}
        return {
            'type': 'ir.actions.act_window',
            'res_model': 'hr.shift.schedule',
            'res_id': schedule.id,
            'view_mode': 'form',
            'target': 'current',
        }
