# -*- coding: utf-8 -*-
import base64
import io
from datetime import date

import pytz

from odoo import models, fields, api
from odoo.exceptions import UserError

_BOGOTA_TZ = pytz.timezone('America/Bogota')


def _bog_hm(dt):
    """Convierte un datetime UTC (naive) a hora Bogotá y retorna string HH:MM."""
    if not dt:
        return ''
    return dt.replace(tzinfo=pytz.utc).astimezone(_BOGOTA_TZ).strftime('%H:%M')


class HrBiometricOvertimeReportWizard(models.TransientModel):
    """Wizard para generar el reporte consolidado de horas extras biométricas.

    Genera un PDF (QWeb) y/o Excel con el detalle de horas trabajadas,
    horas extras diurnas/nocturnas, recargos y festivos/dominicales
    para el período seleccionado.
    """
    _name = 'hr.biometric.overtime.report.wizard'
    _description = 'Reporte Consolidado Horas Extras Biométrico'

    # ─── Filtros ──────────────────────────────────────────────────────────
    date_start = fields.Date(
        string='Fecha Inicio',
        required=True,
        default=lambda self: date.today().replace(day=1),
    )
    date_end = fields.Date(
        string='Fecha Fin',
        required=True,
        default=fields.Date.today,
    )
    employee_ids = fields.Many2many(
        'hr.employee',
        string='Empleados',
        help='Dejar vacío para incluir todos.',
    )
    department_id = fields.Many2one(
        'hr.department',
        string='Departamento',
    )
    address_id = fields.Many2one(
        'res.partner',
        string='Sede / Centro de trabajo',
        domain="[('is_company', '=', True)]",
    )
    state = fields.Selection([
        ('all', 'Todos los estados'),
        ('ok', 'Correctos'),
        ('extra_hours', 'Horas Extras'),
        ('late_minor', 'Tardanza leve'),
        ('late_major', 'Tardanza mayor'),
        ('partial_absence', 'Ausencia parcial'),
        ('total_absence', 'Ausencia total'),
        ('biometric_issue', 'Inconsistencia biométrica'),
    ], string='Resultado asistencia', default='all')

    def _get_records(self):
        """Retorna los registros hr.attendance.record filtrados."""
        domain = [
            ('work_date', '>=', self.date_start),
            ('work_date', '<=', self.date_end),
        ]
        # Filtrar por empleados seleccionados o derivados de departamento/sede
        employee_ids = self.employee_ids
        if self.department_id or self.address_id:
            emp_domain = [('active', '=', True)]
            if self.department_id:
                emp_domain.append(('department_id', '=', self.department_id.id))
            if self.address_id:
                emp_domain.append(('address_id', '=', self.address_id.id))
            dept_employees = self.env['hr.employee'].search(emp_domain)
            if employee_ids:
                employee_ids = employee_ids & dept_employees
            else:
                employee_ids = dept_employees
        if self.employee_ids or self.department_id or self.address_id:
            domain.append(('employee_id', 'in', employee_ids.ids))
        if self.state != 'all':
            domain.append(('attendance_result', '=', self.state))
        records = self.env['hr.attendance.record'].search(domain, order='work_date, employee_id, id')
        # Deduplicar: si el mismo empleado+fecha aparece en varias programaciones,
        # conservar solo el registro con más datos (preferir el que tiene check-in real).
        seen = {}
        for rec in records:
            key = (rec.employee_id.id, rec.work_date)
            if key not in seen:
                seen[key] = rec
            elif rec.actual_check_in and not seen[key].actual_check_in:
                seen[key] = rec
        return self.env['hr.attendance.record'].browse([r.id for r in seen.values()])

    def action_print_pdf(self):
        """Genera el reporte en PDF."""
        self.ensure_one()
        records = self._get_records()
        if not records:
            raise UserError('No se encontraron registros para los filtros seleccionados.')
        return self.env.ref(
            'hr_payroll_co_biometric.action_report_biometric_overtime'
        ).report_action(self, data={'wizard_id': self.id})

    def action_export_excel(self):
        """Genera el reporte en Excel (.xlsx) y lo retorna como descarga."""
        self.ensure_one()
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        except ImportError:
            raise UserError('El módulo openpyxl no está instalado. Instálelo con: pip install openpyxl')

        records = self._get_records()
        if not records:
            raise UserError('No se encontraron registros para los filtros seleccionados.')

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Horas Extras'

        # ── Estilos ──────────────────────────────────────────────────────
        header_font = Font(bold=True, color='FFFFFF', size=10)
        header_fill = PatternFill(start_color='1F4E79', end_color='1F4E79', fill_type='solid')
        sub_header_fill = PatternFill(start_color='2E75B6', end_color='2E75B6', fill_type='solid')
        center = Alignment(horizontal='center', vertical='center', wrap_text=True)
        thin_border = Border(
            left=Side(style='thin'), right=Side(style='thin'),
            top=Side(style='thin'), bottom=Side(style='thin'),
        )
        alt_fill = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')

        # ── Título ───────────────────────────────────────────────────────
        ws.merge_cells('A1:R1')
        title_cell = ws['A1']
        title_cell.value = 'REPORTE CONSOLIDADO DE HORAS EXTRAS - BIOMÉTRICO'
        title_cell.font = Font(bold=True, size=14, color='1F4E79')
        title_cell.alignment = center

        ws.merge_cells('A2:R2')
        period_cell = ws['A2']
        period_cell.value = 'Período: %s al %s' % (
            self.date_start.strftime('%d/%m/%Y'),
            self.date_end.strftime('%d/%m/%Y'),
        )
        period_cell.font = Font(italic=True, size=10)
        period_cell.alignment = center

        ws.row_dimensions[3].height = 5  # spacer

        # ── Encabezados de columna ────────────────────────────────────────
        col_headers = [
            ('Fecha', 12),
            ('Empleado', 30),
            ('Cédula', 14),
            ('Turno', 10),
            ('Sede', 18),
            ('Departamento', 20),
            ('Entrada\nPlanificada', 12),
            ('Salida\nPlanificada', 12),
            ('Entrada\nReal', 12),
            ('Salida\nReal', 12),
            ('H. Regulares\nDiurnas', 10),
            ('H. Regulares\nNocturnas', 10),
            ('H.E. Diurnas\n(HED)', 10),
            ('H.E. Nocturnas\n(HEN)', 10),
            ('H.E. Festivo\n(HEDF)', 10),
            ('Festivo', 8),
            ('Dominical', 8),
            ('Total H.E.\nValor ($)', 14),
        ]
        for col_idx, (header, width) in enumerate(col_headers, 1):
            cell = ws.cell(row=4, column=col_idx, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center
            cell.border = thin_border
            ws.column_dimensions[cell.column_letter].width = width

        ws.row_dimensions[4].height = 30

        # ── Datos ──────────────────────────────────────────────────────────
        totals = {
            'day_hours': 0, 'night_hours': 0,
            'extra_day_hours': 0, 'extra_night_hours': 0,
            'holiday_work_hours': 0, 'total_overtime_value': 0,
        }

        for row_idx, rec in enumerate(records, start=5):
            fill = alt_fill if row_idx % 2 == 0 else None
            row_data = [
                rec.work_date.strftime('%d/%m/%Y') if rec.work_date else '',
                rec.employee_id.name or '',
                rec.employee_id.identification_id or '',
                rec.shift_id.code or '',
                (rec.schedule_line_id.schedule_id.address_id.name or '') if rec.schedule_line_id else '',
                rec.employee_id.department_id.name or '',
                _bog_hm(rec.planned_start),
                _bog_hm(rec.planned_end),
                _bog_hm(rec.actual_check_in),
                _bog_hm(rec.actual_check_out),
                rec.day_hours,
                rec.night_hours,
                rec.extra_day_hours,
                rec.extra_night_hours,
                rec.holiday_work_hours,
                'SÍ' if rec.is_holiday else 'NO',
                'SÍ' if rec.is_sunday else 'NO',
                rec.total_overtime_value,
            ]
            for col_idx, value in enumerate(row_data, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.border = thin_border
                if fill:
                    cell.fill = fill
                if col_idx in (1, 15, 16) and isinstance(value, str):
                    cell.alignment = center
                if col_idx >= 11:
                    cell.number_format = '#,##0.00'

            # Acumular totales
            totals['day_hours'] += rec.day_hours or 0
            totals['night_hours'] += rec.night_hours or 0
            totals['extra_day_hours'] += rec.extra_day_hours or 0
            totals['extra_night_hours'] += rec.extra_night_hours or 0
            totals['holiday_work_hours'] += rec.holiday_work_hours or 0
            totals['total_overtime_value'] += rec.total_overtime_value or 0

        # ── Fila de totales ───────────────────────────────────────────────
        total_row = len(records) + 5
        ws.cell(row=total_row, column=1, value='TOTALES').font = Font(bold=True)
        ws.cell(row=total_row, column=2, value=len(records))
        for col_idx, key in enumerate(
            ['day_hours', 'night_hours', 'extra_day_hours', 'extra_night_hours', 'holiday_work_hours'],
            start=11
        ):
            cell = ws.cell(row=total_row, column=col_idx, value=totals[key])
            cell.font = Font(bold=True)
            cell.fill = sub_header_fill
            cell.font = Font(bold=True, color='FFFFFF')
            cell.number_format = '#,##0.00'
            cell.border = thin_border

        total_val_cell = ws.cell(row=total_row, column=18, value=totals['total_overtime_value'])
        total_val_cell.font = Font(bold=True, color='FFFFFF')
        total_val_cell.fill = sub_header_fill
        total_val_cell.number_format = '#,##0.00'
        total_val_cell.border = thin_border

        # Freeze panes
        ws.freeze_panes = 'A5'

        # Guardar en buffer y retornar como descarga
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        file_data = base64.b64encode(buffer.read()).decode()

        filename = 'horas_extras_%s_%s.xlsx' % (
            self.date_start.strftime('%Y%m%d'),
            self.date_end.strftime('%Y%m%d'),
        )
        attachment = self.env['ir.attachment'].create({
            'name': filename,
            'type': 'binary',
            'datas': file_data,
            'mimetype': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        })
        return {
            'type': 'ir.actions.act_url',
            'url': '/web/content/%d?download=true' % attachment.id,
            'target': 'new',
        }

    def get_report_data(self):
        """Retorna los datos para el template QWeb del PDF."""
        self.ensure_one()
        records = self._get_records()
        total_day = sum(r.day_hours for r in records)
        total_night = sum(r.night_hours for r in records)
        total_extra_day = sum(r.extra_day_hours for r in records)
        total_extra_night = sum(r.extra_night_hours for r in records)
        total_holiday = sum(r.holiday_work_hours for r in records)
        total_value = sum(r.total_overtime_value for r in records)
        return {
            'wizard': self,
            'records': records,
            'fmt_time': _bog_hm,
            'total_day': total_day,
            'total_night': total_night,
            'total_extra_day': total_extra_day,
            'total_extra_night': total_extra_night,
            'total_holiday': total_holiday,
            'total_value': total_value,
        }
