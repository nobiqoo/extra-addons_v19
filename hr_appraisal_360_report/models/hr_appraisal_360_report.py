# -*- coding: utf-8 -*-
# Part of Odoo. See LICENSE file for full copyright and licensing details.

import logging
import io
import base64
from odoo import api, fields, models, _
from odoo.exceptions import AccessError, UserError
from odoo.tools import date_utils

_logger = logging.getLogger(__name__)


class HrAppraisal360Report(models.Model):
    """Modelo para generar y gestionar informes 360 de evaluaciones."""
    _name = 'hr.appraisal.360.report'
    _description = 'Informe de retroalimentaciones 360'
    _order = 'create_date desc'

    name = fields.Char(string='Nombre del informe', readonly=True)
    state = fields.Selection(
        [('draft', 'Borrador'), ('generated', 'Generado'), ('exported', 'Exportado')],
        default='draft', readonly=True)
    report_date = fields.Datetime(string='Fecha de generación', readonly=True, default=fields.Datetime.now)
    all_employee = fields.Boolean(string='Todos los empleados', default=True)
    employee_id = fields.Many2one('hr.employee', string='Persona evaluada')
    survey_id = fields.Many2one(
        'survey.survey',
        string='Encuesta de retroalimentación',
        domain=[('survey_type', '=', 'appraisal')])
    date_from = fields.Date(string='Fecha desde')
    date_to = fields.Date(string='Fecha hasta')
    evaluator_id = fields.Many2one('hr.employee', string='Evaluador')
    appraisal_state = fields.Selection(
        [('1_new', 'Borrador'),
         ('2_pending', 'En curso'),
         ('3_done', 'Realizada')],
        string='Estado de evaluación')
    report_lines = fields.One2many('hr.appraisal.360.report.line', 'report_id', string='Líneas del informe', readonly=True)
    total_lines = fields.Integer(string='Total de líneas', compute='_compute_total_lines')
    excel_file = fields.Binary(string='Archivo Excel', readonly=True, attachment=True)
    excel_filename = fields.Char(string='Nombre del archivo Excel', readonly=True)

    @api.depends('report_lines')
    def _compute_total_lines(self):
        for record in self:
            record.total_lines = len(record.report_lines)

    @api.model_create_multi
    def create(self, vals_list):
        for vals in vals_list:
            if not vals.get('name'):
                vals['name'] = _('Informe 360 - %s') % fields.Datetime.now()
        return super().create(vals_list)

    def action_generate_report(self):
        self.ensure_one()
        if not self.env.user.has_group('hr_appraisal.group_hr_appraisal_user'):
            raise AccessError(_('No tiene permisos para acceder a los informes de evaluaciones.'))

        self.report_lines.unlink()
        domain = self._get_appraisal_domain()
        appraisals = self.env['hr.appraisal'].search(domain)
        if not appraisals:
            raise UserError(_('No se encontraron evaluaciones con los filtros indicados.'))

        lines_to_create = self._prepare_survey_feedback_lines(appraisals)

        if not lines_to_create:
            raise UserError(_('No se encontraron respuestas de encuestas de retroalimentación con los filtros indicados.'))

        self.env['hr.appraisal.360.report.line'].create(lines_to_create)
        self.state = 'generated'
        return {
            'type': 'ir.actions.client',
            'tag': 'display_notification',
            'params': {
                'title': _('Informe generado'),
                'message': _('El informe se generó correctamente con %d líneas.') % len(lines_to_create),
                'type': 'success',
                'sticky': False,
            }
        }

    def _get_appraisal_domain(self):
        domain = []
        if not self.all_employee and self.employee_id:
            domain.append(('employee_id', '=', self.employee_id.id))
        if self.appraisal_state:
            domain.append(('state', '=', self.appraisal_state))
        return domain

    def _get_user_input_domain(self, appraisals):
        domain = [
            ('appraisal_id', 'in', appraisals.ids),
            ('survey_id.survey_type', '=', 'appraisal'),
        ]
        if self.survey_id:
            domain.append(('survey_id', '=', self.survey_id.id))
        if self.date_from:
            domain.append(('end_datetime', '>=', fields.Datetime.to_datetime(self.date_from)))
        if self.date_to:
            domain.append(('end_datetime', '<=', date_utils.end_of(fields.Datetime.to_datetime(self.date_to), 'day')))
        if self.evaluator_id:
            partner_ids = (self.evaluator_id.work_contact_id | self.evaluator_id.user_id.partner_id).ids
            emails = [email for email in (self.evaluator_id.work_email, self.evaluator_id.user_id.email) if email]
            evaluator_domain = []
            if partner_ids:
                evaluator_domain = ['|', ('partner_id', 'in', partner_ids), ('email', 'in', emails or [''])]
            elif emails:
                evaluator_domain = [('email', 'in', emails)]
            else:
                evaluator_domain = [('id', '=', 0)]
            domain += evaluator_domain
        return domain

    def _prepare_survey_feedback_lines(self, appraisals):
        lines_to_create = []
        user_inputs = self.env['survey.user_input'].search(self._get_user_input_domain(appraisals))
        if not user_inputs:
            return lines_to_create

        answers = self.env['survey.user_input.line'].search([
            ('user_input_id', 'in', user_inputs.ids),
            ('answer_type', '!=', False),
            ('question_id.is_page', '=', False),
            ('skipped', '=', False),
        ])
        employees_by_partner, employees_by_email = self._get_evaluator_maps(user_inputs)
        for appraisal in appraisals:
            appraisal_user_inputs = user_inputs.filtered(lambda user_input: user_input.appraisal_id == appraisal)
            for user_input in appraisal_user_inputs:
                evaluator = self._get_evaluator_name(user_input, employees_by_partner, employees_by_email)
                user_input_answers = answers.filtered(lambda line: line.user_input_id == user_input)
                if not user_input_answers:
                    lines_to_create.append(self._prepare_feedback_line_values(appraisal, user_input, evaluator))
                    continue
                for answer_line in user_input_answers:
                    lines_to_create.append(self._prepare_feedback_line_values(appraisal, user_input, evaluator, answer_line))
        return lines_to_create

    def _prepare_feedback_line_values(self, appraisal, user_input, evaluator, answer_line=False):
        values = {
            'report_id': self.id,
            'appraisal_id': appraisal.id,
            'survey_id': user_input.survey_id.id,
            'survey_user_input_id': user_input.id,
            'evaluated_employee_id': appraisal.employee_id.id,
            'evaluated_employee': appraisal.employee_id.display_name,
            'evaluator_employee_id': evaluator.id if evaluator else False,
            'evaluator': evaluator.display_name if evaluator else self._get_user_input_participant_name(user_input),
            'feedback_survey': user_input.survey_id.display_name,
            'survey_response_state': user_input.state,
            'appraisal_date': fields.Date.to_date(user_input.end_datetime) or appraisal.date_close,
            'appraisal_state': appraisal.state,
        }
        if answer_line:
            values.update({
                'survey_user_input_line_id': answer_line.id,
                'question': answer_line.question_id.title,
                'answer': self._get_answer_text(answer_line),
            })
        return values

    def _get_evaluator_maps(self, user_inputs):
        partner_ids = user_inputs.mapped('partner_id').ids
        emails = [email for email in user_inputs.mapped('email') if email]
        employees = self.env['hr.employee'].search([
            '|',
            '|', ('work_contact_id', 'in', partner_ids),
            ('user_id.partner_id', 'in', partner_ids),
            '|', ('work_email', 'in', emails or ['']),
            ('user_id.email', 'in', emails or ['']),
        ])
        employees_by_partner = {}
        employees_by_email = {}
        for employee in employees:
            for partner in employee.work_contact_id | employee.user_id.partner_id:
                employees_by_partner[partner.id] = employee
            for email in (employee.work_email, employee.user_id.email):
                if email:
                    employees_by_email[email.lower()] = employee
        return employees_by_partner, employees_by_email

    def _get_evaluator_name(self, user_input, employees_by_partner, employees_by_email):
        if user_input.partner_id and user_input.partner_id.id in employees_by_partner:
            return employees_by_partner[user_input.partner_id.id]
        if user_input.email and user_input.email.lower() in employees_by_email:
            return employees_by_email[user_input.email.lower()]
        return self.env['hr.employee']

    def _get_user_input_participant_name(self, user_input):
        return user_input.partner_id.display_name or user_input.email or user_input.nickname or _('Evaluador desconocido')

    def _get_answer_text(self, answer_line):
        if answer_line.answer_type == 'suggestion':
            if answer_line.matrix_row_id:
                return '%s: %s' % (answer_line.matrix_row_id.value, answer_line.suggested_answer_id.value)
            return answer_line.suggested_answer_id.value
        answer = answer_line._get_answer_value()
        if answer is False or answer is None:
            return ''
        return str(answer)

    def action_export_excel(self):
        self.ensure_one()
        if self.state == 'draft':
            raise UserError(_('Primero debe generar el informe antes de exportarlo.'))
        excel_content = self._generate_excel_file()
        if excel_content:
            filename = f"informe_retroalimentacion_360_{fields.Date.today().isoformat()}.xlsx"
            self.write({
                'excel_file': excel_content,
                'excel_filename': filename,
                'state': 'exported'
            })
            return {
                'type': 'ir.actions.act_url',
                'url': f'/web/content/{self._name}/{self.id}/excel_file/{filename}?download=true',
                'target': 'new',
            }
        raise UserError(_('No se pudo generar el archivo Excel.'))

    def _generate_excel_file(self):
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        except ImportError:
            raise UserError(_('La librería openpyxl es requerida para exportar a Excel. Por favor instálela.'))
        wb = Workbook()
        ws = wb.active
        ws.title = 'Evaluacion 360'
        header_fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
        header_font = Font(bold=True, color='FFFFFF')
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        headers = [
            _('Persona evaluada'),
            _('Evaluador'),
            _('Pregunta'),
            _('Respuesta'),
            _('Encuesta de retroalimentación'),
            _('Fecha'),
            _('Estado')
        ]
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num)
            cell.value = header
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            cell.border = border
        for row_num, line in enumerate(self.report_lines, 2):
            data = [
                line.evaluated_employee,
                line.evaluator,
                line.question,
                line.answer,
                line.feedback_survey,
                line.appraisal_date,
                line._get_selection_label('survey_response_state'),
            ]
            for col_num, value in enumerate(data, 1):
                cell = ws.cell(row=row_num, column=col_num)
                cell.value = value
                cell.alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)
                cell.border = border
        ws.column_dimensions['A'].width = 25
        ws.column_dimensions['B'].width = 25
        ws.column_dimensions['C'].width = 30
        ws.column_dimensions['D'].width = 30
        ws.column_dimensions['E'].width = 30
        ws.column_dimensions['F'].width = 15
        ws.column_dimensions['G'].width = 15
        ws.column_dimensions['H'].width = 15
        summary_ws = wb.create_sheet('Resumen')
        summary_ws['A1'] = _('Resumen del informe de retroalimentaciones 360')
        summary_ws['A2'] = _('Fecha del informe:')
        summary_ws['B2'] = self.report_date
        summary_ws['A3'] = _('Total de líneas:')
        summary_ws['B3'] = self.total_lines
        summary_ws['A4'] = _('Generado por:')
        summary_ws['B4'] = self.env.user.name
        excel_bytes = io.BytesIO()
        wb.save(excel_bytes)
        excel_bytes.seek(0)
        return base64.b64encode(excel_bytes.read())


class HrAppraisal360ReportLine(models.Model):
    """Líneas del informe de retroalimentaciones 360."""
    _name = 'hr.appraisal.360.report.line'
    _description = 'Línea del informe de retroalimentaciones 360'
    _order = 'id'

    report_id = fields.Many2one('hr.appraisal.360.report', string='Informe', ondelete='cascade', required=True)
    appraisal_id = fields.Many2one('hr.appraisal', string='Evaluación', readonly=True)
    survey_id = fields.Many2one('survey.survey', string='Encuesta', readonly=True)
    survey_user_input_id = fields.Many2one('survey.user_input', string='Respuesta de encuesta', readonly=True)
    survey_user_input_line_id = fields.Many2one('survey.user_input.line', string='Línea de respuesta de encuesta', readonly=True)
    evaluated_employee_id = fields.Many2one('hr.employee', string='Persona evaluada', readonly=True)
    evaluated_employee = fields.Char(string='Persona evaluada', readonly=True)
    evaluator_employee_id = fields.Many2one('hr.employee', string='Evaluador', readonly=True)
    evaluator = fields.Char(string='Evaluador', readonly=True)
    question = fields.Text(string='Pregunta', readonly=True)
    answer = fields.Text(string='Respuesta', readonly=True)
    feedback_survey = fields.Char(string='Encuesta de retroalimentación', readonly=True)
    survey_response_state = fields.Selection(
        [('new', 'No empezada'),
         ('in_progress', 'En progreso'),
         ('done', 'Terminada')],
        string='Estado', readonly=True)
    appraisal_date = fields.Date(string='Fecha', readonly=True)
    appraisal_state = fields.Selection(
        [('1_new', 'Borrador'),
         ('2_pending', 'En curso'),
         ('3_done', 'Realizada')],
        string='Estado de evaluación', readonly=True)

    def _get_selection_label(self, field_name):
        self.ensure_one()
        value = self[field_name]
        if not value:
            return ''
        return dict(self._fields[field_name].selection).get(value, value)
